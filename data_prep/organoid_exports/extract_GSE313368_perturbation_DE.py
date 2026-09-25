#!/usr/bin/env python3
"""Stream the published GSE313368 cell-type DE workbook into one auditable TSV."""

from __future__ import annotations

import csv
import gzip
import hashlib
import os
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def as_text(value: object) -> str:
    return "" if value is None else str(value)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: extract_GSE313368_perturbation_DE.py <project_dir>")

    root = Path(sys.argv[1])
    source = root / "01_data" / "41467_2025_68247_MOESM10_ESM.xlsx"
    output = root / "02_results" / "GSE313368_celltype_perturbation_DE_all.tsv.gz"
    audit_output = root / "tables_for_article" / "Table_S20_GSE313368_perturbation_data_audit.tsv"
    tmp_output = output.with_name(output.name + ".tmp.gz")
    tmp_audit = audit_output.with_name(audit_output.name + ".tmp")

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet_celltypes = {
        "A_TA_DEGs": "Transit_Amplifying",
        "B_Stem_DEGs": "Stem",
        "C_Goblet_DEGs": "Goblet",
        "D_Colonocyte_DEGs": "Colonocyte",
    }
    expected_header = ["gene", "baseMean", "log2FoldChange", "lfcSE", "pvalue", "padj", "pert"]
    hashes: dict[str, str] = {}
    audit_rows: list[dict[str, object]] = []

    with gzip.open(tmp_output, "wt", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow([
            "dataset", "source_workbook", "source_sheet", "celltype", "gene",
            "baseMean", "log2FoldChange", "lfcSE", "pvalue", "padj", "perturbation",
            "supplement_qc_exclude_gene_level",
        ])
        for sheet_name, celltype in sheet_celltypes.items():
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            header = [as_text(v) for v in next(rows)]
            if header != expected_header:
                raise RuntimeError(f"Unexpected header in {sheet_name}: {header}")

            digest = hashlib.sha256()
            n_rows = 0
            perturbations: Counter[str] = Counter()
            for row in rows:
                values = [as_text(v) for v in row]
                digest.update(("\t".join(values) + "\n").encode("utf-8"))
                gene, base_mean, log_fc, lfc_se, pvalue, padj, perturbation = values
                if not gene or not perturbation:
                    continue
                n_rows += 1
                perturbations[perturbation] += 1
                # The publisher workbook contains byte-identical Stem and Goblet sheets.
                # Retain them for provenance but exclude both from gene-level inference.
                exclude = sheet_name in {"B_Stem_DEGs", "C_Goblet_DEGs"}
                writer.writerow([
                    "GSE313368", source.name, sheet_name, celltype, gene, base_mean,
                    log_fc, lfc_se, pvalue, padj, perturbation, str(exclude).upper(),
                ])
            hashes[sheet_name] = digest.hexdigest()
            audit_rows.append({
                "dataset": "GSE313368",
                "source_file": source.name,
                "source_sheet": sheet_name,
                "data_level": "celltype_perturbation_DE",
                "celltype": celltype,
                "n_rows": n_rows,
                "n_perturbations": len(perturbations),
                "sha256_of_tab_separated_data_rows": hashes[sheet_name],
                "qc_issue": "",
                "analysis_use": "gene_level_module_screen",
            })

    workbook.close()
    duplicate = hashes["B_Stem_DEGs"] == hashes["C_Goblet_DEGs"]
    if not duplicate:
        raise RuntimeError("Expected publisher duplicate-sheet audit did not reproduce")
    for row in audit_rows:
        if row["source_sheet"] in {"B_Stem_DEGs", "C_Goblet_DEGs"}:
            row["qc_issue"] = "publisher_workbook_stem_and_goblet_sheets_are_byte_identical"
            row["analysis_use"] = "provenance_only_excluded_from_gene_level_inference"

    audit_fields = list(audit_rows[0])
    with open(tmp_audit, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=audit_fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(audit_rows)

    os.replace(tmp_output, output)
    os.replace(tmp_audit, audit_output)
    print(f"Wrote {output}")
    print(f"Wrote {audit_output}")
    print("QC: publisher Stem and Goblet DE sheets are byte-identical; both flagged and excluded.")


if __name__ == "__main__":
    main()
