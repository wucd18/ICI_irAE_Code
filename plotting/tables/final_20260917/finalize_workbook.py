from common import *
from lxml import etree as E
from copy import deepcopy
import openpyxl
X='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ns={'s':X}
q=lambda s:'{'+X+'}'+s
src=ZIN/'additional/Additional_file_2.xlsx';out=U/'Additional_file_2.xlsx'
edits=jr(A/'artifact_authored_cell_changes.json')
wb=openpyxl.load_workbook(src);auth=openpyxl.load_workbook(C/'artifact_authored_intermediate.xlsx')
with zipfile.ZipFile(src) as z:parts={n:z.read(n) for n in z.namelist()}
styles=E.fromstring(parts['xl/styles.xml']);xfs=styles.find(q('cellXfs'));newstyles={};formats=[]
for sheet in wb.sheetnames:
 path='xl/worksheets/sheet'+str(wb.sheetnames.index(sheet)+1)+'.xml';root=E.fromstring(parts[path]);ws=wb[sheet]
 for e in [t for t in edits if t['sheet']==sheet]:
  assert auth[sheet][e['cell']].value==e['new_value']
  cell=root.find('.//s:c[@r="'+e['cell']+'"]',ns);assert cell is not None
  assert ws[e['cell']].value==e['old_value']
  cell.set('t','inlineStr')
  for child in list(cell):
   if child.tag in [q('v'),q('is')]:cell.remove(child)
  t=E.SubElement(E.SubElement(cell,q('is')),q('t'));t.text=e['new_value'];t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
  if e['new_value'].startswith('Source:'):
   sid=cell.get('s','0')
   if sid not in newstyles:
    xf=deepcopy(xfs[int(sid)]);al=xf.find(q('alignment'))
    if al is None:al=E.SubElement(xf,q('alignment'))
    al.set('wrapText','1');xf.set('applyAlignment','1');newstyles[sid]=str(len(xfs));xfs.append(xf)
   cell.set('s',newstyles[sid]);row=cell.getparent();oldht=row.get('ht');row.set('ht',str(max(float(oldht or 15),42)));row.set('customHeight','1')
   formats.append(dict(sheet=sheet,cell=e['cell'],old_height=oldht,new_height=row.get('ht'),reason='DOI source note wrapping only'))
 parts[path]=E.tostring(root)
xfs.set('count',str(len(xfs)));parts['xl/styles.xml']=E.tostring(styles)
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
after=openpyxl.load_workbook(out);allowed={(e['sheet'],e['cell']):e for e in edits};checks=[]
for s in wb.sheetnames:
 a,b=wb[s],after[s];assert a.max_row==b.max_row and a.max_column==b.max_column
 assert str(a.merged_cells)==str(b.merged_cells)
 for row in a:
  for c in row:
   d=b[c.coordinate];key=(s,c.coordinate);expected=allowed[key]['new_value'] if key in allowed else c.value
   assert d.value==expected,(key,c.value,d.value)
   assert c.number_format==d.number_format
   if c.value is not None:checks.append(dict(sheet=s,cell=c.coordinate,data_type=c.data_type,status='PASS',authorized_text_change=key in allowed))
# XML cells: every non-whitelist cell (including f/v cached values) is byte-canonical identical.
with zipfile.ZipFile(src) as za,zipfile.ZipFile(out) as zb:
 for i,s in enumerate(wb.sheetnames,1):
  na='xl/worksheets/sheet'+str(i)+'.xml';a=E.fromstring(za.read(na));b=E.fromstring(zb.read(na))
  for x,y in zip(a.findall('.//s:c',ns),b.findall('.//s:c',ns)):
   assert x.get('r')==y.get('r')
   if (s,x.get('r')) not in allowed:assert E.tostring(x)==E.tostring(y),(s,x.get('r'))
  for tag in ['mergeCells','sheetViews','cols','pageMargins','pageSetup','printOptions','autoFilter','hyperlinks']:
   xa,xb=a.find(q(tag)),b.find(q(tag));assert (E.tostring(xa) if xa is not None else None)==(E.tostring(xb) if xb is not None else None),(s,tag)
 for n in za.namelist():
  if not re.match(r'xl/worksheets/sheet\d+.xml$',n) and n!='xl/styles.xml':assert za.read(n)==zb.read(n),n
cw(A/'XLSX_CELL_DIFF.csv',edits);cw(A/'XLSX_CELL_INVARIANCE.csv',checks);cw(A/'XLSX_NOTE_FORMAT_CHANGES.csv',formats)
jw(A/'XLSX_INTEGRITY.json',dict(status='PASS',all_populated_cells_checked=len(checks),authorized_text_edits=len(edits),formula_and_cached_values_unchanged=True,unrelated_XML_parts_unchanged=True,source_sha256=sha(src),output_sha256=sha(out)))
print('XLSX PASS',len(edits),len(checks))
